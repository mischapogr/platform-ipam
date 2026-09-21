#!/usr/bin/env python3
"""Generate .xlsx fixtures for platform-ipam internal/onboard D2 verification
using openpyxl. Run inside python:3.13-alpine with `pip install openpyxl`.

Produces (in the current directory):
  openpyxl_main.xlsx          - Networks sheet + Accounts sheet, the full
                                 "rows covering" scenario list from the D2
                                 work order, plus a sparse (omitted) blank
                                 row in the middle to test row-number fidelity.
  openpyxl_merged_header.xlsx - a merged header cell (D1:E1), value only in
                                 the top-left cell of the merge.
  openpyxl_styled_date.xlsx   - a date-formatted/styled cell in an extra
                                 "onboarded" column.
  openpyxl_header_offset.xlsx - header row is not row 1: rows 1-2 are never
                                 written at all (the realistic case: nobody
                                 typed anything there), so they are entirely
                                 absent from the worksheet XML.
"""
import openpyxl
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

NBSP = " "

HEADER = ["Account ID", "Region", "CIDR", "VPC ID", "Name"]


def write_main_rows(ws, start_row=1, skip_row8=True):
    for col, h in enumerate(HEADER, start=1):
        ws.cell(row=start_row, column=col, value=h)  # row start_row
    r = start_row + 1
    ws.cell(row=r, column=1, value="123456789012")
    ws.cell(row=r, column=2, value="us-east-1")
    ws.cell(row=r, column=3, value="10.0.0.0/16")
    ws.cell(row=r, column=4, value="vpc-0123abcd0000")
    ws.cell(row=r, column=5, value="prod-main")
    r += 1
    ws.cell(row=r, column=1, value=12345678)  # NUMBER, lost leading zeros
    ws.cell(row=r, column=2, value="us-west-2")
    ws.cell(row=r, column=3, value="10.1.0.0/16")
    ws.cell(row=r, column=4, value="vpc-0456def00000")
    ws.cell(row=r, column=5, value="staging")
    r += 1
    ws.cell(row=r, column=1, value="000123456789")  # TEXT, leading zeros kept
    ws.cell(row=r, column=2, value="eu-west-1")
    ws.cell(row=r, column=3, value="10.2.0.0/16")
    ws.cell(row=r, column=4, value="vpc-0789ghi00000")
    ws.cell(row=r, column=5, value="dev")
    r += 1
    ws.cell(row=r, column=1, value=123456789012)  # NUMBER, 12-digit
    ws.cell(row=r, column=2, value="ap-southeast-1")
    ws.cell(row=r, column=3, value="10.3.0.0/16")
    ws.cell(row=r, column=4, value="vpc-0abc12340000")
    ws.cell(row=r, column=5, value="qa")
    r += 1
    ws.cell(row=r, column=1, value="234567890123")
    ws.cell(row=r, column=2, value="sa-east-1")
    c = ws.cell(row=r, column=3, value="10.4.0.0/16\n10.5.0.0/16")  # 2 CIDRs, line break
    c.alignment = Alignment(wrap_text=True)
    ws.cell(row=r, column=4, value="vpc-0def56780000")
    ws.cell(row=r, column=5, value="multi-cidr")
    r += 1
    ws.cell(row=r, column=1, value="345678901234")
    ws.cell(row=r, column=2, value="ca-central-1")
    ws.cell(row=r, column=3, value="10.6.0.0/16")
    ws.cell(row=r, column=4, value="vpc-0ghi90120000")
    ws.cell(row=r, column=5, value=f"Acme,{NBSP}Inc")  # comma + NBSP
    r += 1
    if skip_row8:
        r += 1  # leave this spreadsheet row entirely untouched (real writers omit it)
    ws.cell(row=r, column=1, value="456789012345")
    ws.cell(row=r, column=2, value="af-south-1")
    ws.cell(row=r, column=3, value="10.7.0.0/16")
    ws.cell(row=r, column=4, value="vpc-0jkl34560000")
    ws.cell(row=r, column=5, value="after-gap")
    return r  # last written row number


def main_workbook():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Networks"
    write_main_rows(ws)

    ws2 = wb.create_sheet("Accounts")
    ws2.append(["Account ID", "Account Name"])
    ws2.cell(row=2, column=1, value="111111111111")
    ws2.cell(row=2, column=2, value="Example Co")
    ws2.cell(row=3, column=1, value=22222222)  # NUMBER, lost leading zeros
    ws2.cell(row=3, column=2, value="Other Co")

    wb.save("openpyxl_main.xlsx")


def merged_header_workbook():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Networks"
    ws.append(["Account ID", "Region", "CIDR", "VPC ID", "Name"])
    ws.merge_cells("D1:E1")
    ws["D1"] = "VPC ID / Name"  # merged header: value only in the top-left cell
    ws.cell(row=2, column=1, value="123456789012")
    ws.cell(row=2, column=2, value="us-east-1")
    ws.cell(row=2, column=3, value="10.0.0.0/16")
    ws.cell(row=2, column=4, value="vpc-0123abcd0000")
    ws.cell(row=2, column=5, value="prod-main")
    wb.save("openpyxl_merged_header.xlsx")


def styled_date_workbook():
    import datetime
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Networks"
    ws.append(["Account ID", "Region", "CIDR", "VPC ID", "Name", "Onboarded"])
    c1 = ws.cell(row=2, column=1, value="123456789012")
    ws.cell(row=2, column=2, value="us-east-1")
    ws.cell(row=2, column=3, value="10.0.0.0/16")
    ws.cell(row=2, column=4, value="vpc-0123abcd0000")
    ws.cell(row=2, column=5, value="prod-main")
    d = ws.cell(row=2, column=6, value=datetime.date(2024, 1, 15))
    d.number_format = "MM/DD/YYYY"
    # also style the account id cell (bold + fill) to prove a styled cell's
    # value is read the same as an unstyled one
    from openpyxl.styles import Font, PatternFill
    c1.font = Font(bold=True)
    c1.fill = PatternFill("solid", fgColor="FFFF00")
    wb.save("openpyxl_styled_date.xlsx")


def header_offset_workbook():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Networks"
    # Rows 1-2 are never touched at all: real spreadsheet row numbers 1 and 2
    # simply do not exist as <row> elements in the sheet XML.
    write_main_rows(ws, start_row=3, skip_row8=False)
    wb.save("openpyxl_header_offset.xlsx")


if __name__ == "__main__":
    main_workbook()
    merged_header_workbook()
    styled_date_workbook()
    header_offset_workbook()
    print("openpyxl version:", openpyxl.__version__)
